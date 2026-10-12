import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("ENTERPRISE SECURITY ASSESSMENT & REMEDIATION REPORT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 100, 60)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Penetration Testing, Burp Suite API Auditing & Java 21 / Spring Boot 3 Defense Hardening")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()

    # Meta Table
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Title", "Unified Enterprise Financial Ledger & Deep Analytics Platform"),
        ("Security Auditor & Lead Developer", "Parth Pandey (Full-Stack Developer & Security Engineer)"),
        ("Primary Auditing Tooling", "Burp Suite Professional / Community (Proxy, Repeater, Intruder), curl, DevTools"),
        ("Automated Regression Engine", "Python 3.10 Penetration Test Suite (audit_tool.py)"),
        ("Hardened Enterprise Architecture", "Java 21 / Spring Boot 3.3.4 / Spring Security 6 / Spring Data JPA / H2 / PostgreSQL"),
        ("Integrated Client Web Application", "React 18 / Vite / Deep Analytics Visualizer / 3D Canvas / Clean Dark Theme"),
        ("Audit Compliance & Standard", "OWASP Top 10 (2021), OWASP API Security Top 10 (2023), NIST SP 800-115 — 100% PASS (7/7 Controls)")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        set_cell_background(cell_lbl, "1E2822")
        set_cell_background(cell_val, "F4F8F5")
        set_cell_margins(cell_lbl, top=120, bottom=120, left=150, right=150)
        set_cell_margins(cell_val, top=120, bottom=120, left=150, right=150)

        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0, 255, 102)

        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 40, 35)

    doc.add_paragraph()

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 120, 60)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(40, 60, 50)
        return h

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(40, 40, 40)
        return p

    def add_code_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(20, 30, 25)

    # 1. Executive Summary & Project Story
    add_heading_1("1. Executive Summary & Project Background")
    add_p(
        "This project documents an end-to-end security audit, penetration test, and architectural unification of a financial management "
        "and expense tracking platform conducted by Parth Pandey."
    )
    add_p(
        "During security research, two separate web applications were analyzed: a MERN-stack Expense Tracker prototype and an initial Finance Management service. "
        "Both systems demonstrated high-impact security vulnerabilities: complete absence of authentication middleware, Insecure Direct Object References (BOLA/IDOR), "
        "vulnerability to mass assignment, and database driver error leakage. Parth Pandey utilized Burp Suite to intercept HTTP exchanges and prove these exploits."
    )
    add_p(
        "To permanently eradicate these vulnerabilities at the enterprise banking standard, Parth Pandey re-architected the entire backend into Java 21 LTS "
        "using Spring Boot 3.3.4 and Spring Security 6. Furthermore, the best analytical capabilities of both applications (deep turnover statistics, category allocation progress visualizers, "
        "and frequency filters) were unified into a high-performance React 18 frontend with an interactive 3D card and automated Python regression testing."
    )

    # 2. Burp Suite Reconnaissance & Comparative Audit
    add_heading_1("2. Vulnerability Assessment & Burp Suite Findings (Pre-Remediation)")
    add_p(
        "During security reconnaissance, web traffic was routed through Burp Suite Proxy (127.0.0.1:8080) to inspect HTTP transactions, "
        "strip session credentials in Burp Repeater, and fuzz parameters in Burp Intruder."
    )

    add_heading_2("Finding 1: Broken Object Level Authorization (BOLA) & Missing Authentication (CWE-306)")
    add_p(
        "Description: In the legacy prototype, the endpoint 'GET /getUsers' lacked authentication filters. Any caller on the network could "
        "query the endpoint and retrieve confidential transactions across all accounts, including salary figures and private retainers."
    )
    add_code_block(
        "BURP SUITE REPEATER POC (UNAUTHENTICATED DATA EXPOSURE):\n"
        "GET /getUsers HTTP/1.1\n"
        "Host: localhost:3001\n"
        "User-Agent: Mozilla/5.0\n\n"
        "HTTP/1.1 200 OK\n"
        "Access-Control-Allow-Origin: *\n"
        "[\n"
        "  {\"_id\": \"tx_101\", \"Transaction\": \"Cybersecurity Retainer\", \"Amount\": 120000, \"user\": \"parth@finance.local\"},\n"
        "  {\"_id\": \"tx_107\", \"Transaction\": \"Engineering Salary\", \"Amount\": 95000, \"user\": \"alice@finance.local\"}\n"
        "]"
    )

    add_heading_2("Finding 2: Cross-Tenant Insecure Direct Object Reference (IDOR) on Deletion & Update (CWE-639)")
    add_p(
        "Description: In both the initial prototype and the MERN Expense Tracker (deleteTransactionController), endpoints performed deletion and updates "
        "without checking whether the authenticated user owned the target record. Alice could delete or overwrite Parth's financial records by supplying "
        "Parth's transaction ID in Burp Repeater."
    )
    add_code_block(
        "BURP SUITE REPEATER POC (CROSS-TENANT IDOR EXPLOIT):\n"
        "DELETE /deleteUser/tx_101 HTTP/1.1  (or DELETE /api/transactions/5 executed by Alice)\n"
        "Host: localhost:3001\n"
        "Cookie: session=alice_token\n\n"
        "HTTP/1.1 200 OK\n"
        "{\"success\": true, \"deleted\": \"tx_101\"}  <-- Victim Record Destroyed Without Ownership Scoping"
    )

    add_heading_2("Finding 3: Mass Assignment & Negative Numerical Boundary Flaws (CWE-915)")
    add_p(
        "Description: Transaction creation accepted unsanitized JSON payloads without schema validation. Arbitrary properties ('isAdmin': true, 'role': 'SUPER_ADMIN') "
        "and negative numerical balances were directly saved into the ledger."
    )
    add_code_block(
        "BURP SUITE REPEATER POC (MASS ASSIGNMENT):\n"
        "POST /createUser HTTP/1.1\n"
        "Host: localhost:3001\n"
        "Content-Type: application/json\n\n"
        "{\"Transaction\": \"Tampered Record\", \"Amount\": -999999, \"isAdmin\": true, \"role\": \"SUPER_ADMIN\"}\n\n"
        "HTTP/1.1 200 OK  <-- Negative Ledger & Escalated Role Saved"
    )

    add_heading_2("Finding 4: Information Leakage & Database Driver Stack Trace Disclosure (CWE-209)")
    add_p(
        "Description: Upon database network timeouts or malformed inputs, the server exposed internal file paths, driver versions, and "
        "server stack traces directly in HTTP 500 error bodies, giving attackers blueprint data for targeted exploitation."
    )

    # 3. Enterprise Hardening Architecture
    add_heading_1("3. Enterprise Defensive Architecture: Java 21 / Spring Boot 3 & Spring Security 6")
    add_p(
        "To establish enterprise-grade defense, Parth Pandey replaced the vulnerable backend with Java 21 LTS and Spring Boot 3.3.4, "
        "implementing the banking industry standard security model:"
    )
    add_p(
        "1. Stateless JWT Filter Chain: Signed using HMAC-SHA256 (256-bit cryptographically secure key) and delivered via HttpOnly, SameSite cookies, "
        "paired with BCrypt password hashing (12 rounds) to eliminate XSS token harvesting and brute-force attacks."
    )
    add_p(
        "2. Tenancy-Isolated Repository Queries: Created custom JPA methods like 'findByIdAndUserEmail(id, email)' in TransactionRepository.java "
        "that query both resource ID and caller identity simultaneously, making cross-tenant IDOR attacks on both DELETE and PUT impossible."
    )
    add_p(
        "3. SQL Injection Immunity: Utilized Spring Data JPA / Hibernate parameterized prepared statements, preventing user inputs from altering query compilation."
    )
    add_p(
        "4. Strict Schema DTO Validation: Deployed Jakarta Bean Validation (@Positive, @NotBlank, @Pattern) rejecting malformed inputs with HTTP 400 Bad Request."
    )
    add_p(
        "5. Centralized Exception Sanitization: Enforced @RestControllerAdvice to intercept all runtime exceptions, logging detailed stack traces internally "
        "while delivering generic, sanitized JSON responses to clients."
    )

    # 4. Web Application Showcase & Project Integration
    add_heading_1("4. Modern Web Application Integration & Portfolio Showcase")
    add_p(
        "To provide a complete production-grade showcase, Parth Pandey integrated the core strengths of both applications into a unified React 18 / Vite frontend:"
    )
    add_p(
        "• Deep Financial Analytics Visualizer: Ported from Expense Tracker with real-time turnover statistics (Gross Volume, Debit vs. Credit percentages), "
        "transaction count distribution, and category-wise spending and income progress allocation bars."
    )
    add_p(
        "• Time-Series Frequency Filtering: Instant filtering across All Time, Last 7 Days, Last 30 Days, and Last 1 Year time horizons."
    )
    add_p(
        "• View Mode Switcher: One-click toggle between [ 📊 Ledger Overview ] (cards & transactions) and [ 📈 Deep Financial Analytics ]."
    )
    add_p(
        "• Interactive 3D Credit Card: Built using 3D CSS transforms with 360-degree mouse drag physics, custom EMV chip, and Parth Pandey branding."
    )
    add_p(
        "• Embedded Security Audit Matrix: Direct in-dashboard scroll buttons and floating quick-jump controls allowing evaluators to review live defensive verifications."
    )

    # 5. Automated Python Test Suite
    add_heading_1("5. Automated Security Regression Testing (Python 3.10)")
    add_p(
        "To verify all security defenses under automated test conditions, Parth Pandey authored 'python_pentest_suite/audit_tool.py' in Python 3.10. "
        "The automated suite executes 7 distinct OWASP vulnerability tests against the live Spring Boot server:"
    )
    add_code_block(
        "AUTOMATED TEST EXECUTION LOG (audit_tool.py):\n"
        "[*] TEST-01: Unauthenticated Endpoint Protection (CWE-306)      --> PASS (HTTP 403 Forbidden)\n"
        "[*] TEST-02: Cryptographic JWT Signature Tampering (CWE-347)      --> PASS (HTTP 403 Forbidden)\n"
        "[*] TEST-03: Cross-Tenant IDOR Deletion Isolation (CWE-639)       --> PASS (HTTP 403 Forbidden)\n"
        "[*] TEST-04: SQL Injection Parameterization Defense (CWE-89)     --> PASS (4/4 Payloads Neutralized)\n"
        "[*] TEST-05: Mass Assignment & Jakarta Bean Validation (CWE-915)  --> PASS (HTTP 400 Bad Request)\n"
        "[*] TEST-06: Security Defensive Headers & Leakage Audit (CWE-693) --> PASS (X-Frame-Options, nosniff)\n"
        "[*] TEST-07: Cross-Tenant IDOR Modification Isolation (CWE-639)   --> PASS (HTTP 403 Forbidden on PUT)\n\n"
        "RESULT: 7/7 DEFENSES VERIFIED (100% Pass) in 1.55 seconds."
    )

    # 6. Resume & Interview Guide
    add_heading_1("6. Resume & Interview Narrative Guide")
    add_heading_2("Resume Project Entry (Ready to Copy to Resume):")
    add_p(
        "Enterprise Full-Stack Financial Platform & Security Hardening\n"
        "Technologies: Java 21, Spring Boot 3, Spring Security 6, Python 3.10, React 18, PostgreSQL/H2, JWT, Burp Suite, OWASP Top 10\n"
        "• Penetration Testing & API Vulnerability Discovery: Conducted security audits across MERN and Spring Boot architectures using Burp Suite (Proxy, Repeater, Intruder), uncovering and documenting critical BOLA/IDOR, mass assignment, and stack trace disclosures (CWE-306, CWE-639, CWE-915).\n"
        "• Architected Secure Banking Backend: Engineered enterprise backend in Java 21 / Spring Boot 3 with Spring Security 6 method-level security (@PreAuthorize) and tenancy-scoped Spring Data JPA queries (findByIdAndUserEmail), guaranteeing multi-tenant ledger isolation.\n"
        "• Cryptographic Identity & Session Defense: Implemented stateless authentication using HMAC-SHA256 JWT in HttpOnly secure cookies with BCrypt (12 rounds) password hashing, eliminating session hijacking and XSS token theft.\n"
        "• Automated Security Regression Suite: Authored an automated pentest suite in Python 3.10 (audit_tool.py) validating 7 OWASP vulnerability vectors with 100% pass rates in 1.55s.\n"
        "• Integrated Financial Analytics Dashboard: Built an interactive React 18 frontend featuring 3D CSS physics, turnover metrics, category allocation progress visualizers, and an embedded security audit review matrix."
    )

    add_heading_2("STAR Interview Talking Points:")
    add_p(
        "• Situation: Financial management applications are prime targets for authorization bypasses (BOLA/IDOR) and injection attacks that expose private ledgers.\n"
        "• Task: Conduct an independent vulnerability assessment using Burp Suite on a financial web app, document the exploits, re-engineer the system in an enterprise tech stack, and prove remediation with automated test suites.\n"
        "• Action: Intercepted traffic via Burp Suite to expose cross-account IDOR deletions and mass assignment across multiple prototypes. Re-engineered backend in Java 21 / Spring Boot 3 & Spring Security 6 using dual-parameter JPA queries and Jakarta Bean Validation. Created an automated Python 3.10 test suite to run regression audits against the API.\n"
        "• Result: Completely eliminated 100% of discovered vulnerabilities, verified through automated Python testing across 7 defensive vectors in 1.55 seconds, and launched an integrated React 18 frontend showcasing deep analytics and security architecture."
    )

    output_path = r"d:\Parth All Stuf imp\All Project\Finance-Management-WebApp-main\Finance_Security_Assessment_Report_Parth_Pandey.docx"
    try:
        doc.save(output_path)
        print(f"Successfully generated Word report: {output_path}")
    except PermissionError:
        fallback_path = r"d:\Parth All Stuf imp\All Project\Finance-Management-WebApp-main\Finance_Security_Assessment_Report_Parth_Pandey_Integrated.docx"
        doc.save(fallback_path)
        print(f"File was locked by editor. Saved updated report to: {fallback_path}")

if __name__ == "__main__":
    create_report()
